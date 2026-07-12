from collections.abc import Generator
from contextlib import contextmanager
from datetime import timedelta
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app import models  # noqa: F401
from app.database import Base, get_db
from app.main import app
from app.models import Asset, Finding, Project, Scan, ScanAssetResult
from app.models.mixins import utc_now


FUTURE_AUTH_EXPIRY = "2027-07-03T23:59:59Z"


@contextmanager
def _client() -> Generator[tuple[TestClient, sessionmaker[Session]], None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
    Base.metadata.create_all(bind=engine)

    def override_db() -> Generator[Session, None, None]:
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_db
    try:
        yield TestClient(app), TestingSessionLocal
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(bind=engine)


def _auth_headers(client: TestClient) -> dict[str, str]:
    register_response = client.post(
        "/api/v1/auth/register",
        json={"email": "reports@example.org", "full_name": "Report Tester", "password": "StrongPass123!"},
    )
    assert register_response.status_code == 201
    login_response = client.post("/api/v1/auth/login", json={"email": "reports@example.org", "password": "StrongPass123!"})
    assert login_response.status_code == 200
    return {"Authorization": f"Bearer {login_response.json()['access_token']}"}


def _project_id(client: TestClient, headers: dict[str, str], company_name: str = "Report Org") -> str:
    response = client.post(
        "/api/v1/projects",
        headers=headers,
        json={
            "company_name": company_name,
            "main_domain": "example.org",
            "description": "Authorized public scope",
            "scan_frequency": "manual",
            "authorization_confirmed": True,
            "authorization_contact": "owner@example.org",
            "authorization_expires_at": FUTURE_AUTH_EXPIRY,
        },
    )
    assert response.status_code == 201
    return str(response.json()["id"])


def test_excel_report_is_persisted_and_limited_to_selected_scan(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.reports.builder.REPORT_DIR", Path(tmp_path))
    with _client() as (client, session_factory):
        headers = _auth_headers(client)
        project_id = _project_id(client, headers)

        db = session_factory()
        try:
            project = db.get(Project, project_id)
            assert project is not None
            now = utc_now()
            selected_scan = Scan(
                project=project,
                status="completed",
                scan_profile="safe",
                started_at=now - timedelta(minutes=10),
                finished_at=now - timedelta(minutes=8),
                assets_scanned=1,
                findings_created=1,
                risk_score=80,
            )
            other_scan = Scan(
                project=project,
                status="completed",
                scan_profile="aggressive",
                started_at=now - timedelta(minutes=5),
                finished_at=now - timedelta(minutes=3),
                assets_scanned=1,
                findings_created=1,
                risk_score=20,
            )
            db.add_all([selected_scan, other_scan])
            db.flush()
            selected_asset = Asset(
                project_id=project.id,
                scan_id=selected_scan.id,
                hostname="selected.example.org",
                first_seen_at=now,
                last_seen_at=now,
            )
            other_asset = Asset(
                project_id=project.id,
                scan_id=other_scan.id,
                hostname="other.example.org",
                first_seen_at=now,
                last_seen_at=now,
            )
            selected_finding = Finding(
                project_id=project.id,
                scan_id=selected_scan.id,
                asset=selected_asset,
                title="Selected scan finding",
                description="A finding from the selected scan.",
                severity="high",
                category="Exposure",
                first_seen_at=now,
                last_seen_at=now,
            )
            other_finding = Finding(
                project_id=project.id,
                scan_id=other_scan.id,
                asset=other_asset,
                title="Other scan finding",
                description="A finding from a different scan.",
                severity="low",
                category="Headers",
                first_seen_at=now,
                last_seen_at=now,
            )
            db.add_all([selected_asset, other_asset, selected_finding, other_finding])
            db.commit()
            selected_scan_id = selected_scan.id
            other_scan_id = other_scan.id
        finally:
            db.close()

        response = client.post(
            f"/api/v1/projects/{project_id}/reports/excel",
            headers=headers,
            json={"scan_id": selected_scan_id},
        )

        assert response.status_code == 201
        body = response.json()
        assert body["scan_id"] == selected_scan_id
        assert body["scan_profile"] == "safe"
        assert body["scan_date"] is not None

        workbook = load_workbook(body["file_path"])
        assert workbook["Summary"]["B3"].value == selected_scan_id
        assert workbook["Summary"]["B4"].value == "safe"
        assert workbook["Scan Metadata"].max_row == 2
        assert workbook["Scan Metadata"]["A2"].value == selected_scan_id

        asset_hostnames = [row[0] for row in workbook["Assets"].iter_rows(min_row=2, values_only=True)]
        finding_titles = [row[2] for row in workbook["Findings"].iter_rows(min_row=2, values_only=True)]
        assert asset_hostnames == ["selected.example.org"]
        assert finding_titles == ["Selected scan finding"]
        assert "other.example.org" not in asset_hostnames
        assert "Other scan finding" not in finding_titles
        assert other_scan_id != selected_scan_id


def test_report_creation_rejects_scan_from_another_project(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.reports.builder.REPORT_DIR", Path(tmp_path))
    with _client() as (client, session_factory):
        headers = _auth_headers(client)
        project_id = _project_id(client, headers, company_name="Primary Org")
        other_project_id = _project_id(client, headers, company_name="Other Org")

        db = session_factory()
        try:
            other_project = db.get(Project, other_project_id)
            assert other_project is not None
            other_scan = Scan(project=other_project, status="completed", scan_profile="safe")
            db.add(other_scan)
            db.commit()
            other_scan_id = other_scan.id
        finally:
            db.close()

        response = client.post(
            f"/api/v1/projects/{project_id}/reports/pdf",
            headers=headers,
            json={"scan_id": other_scan_id},
        )

        assert response.status_code == 404


def test_historical_report_uses_immutable_manifest_after_canonical_rows_move(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.reports.builder.REPORT_DIR", Path(tmp_path))
    monkeypatch.setattr("app.routes.reports.REPORT_DIR", Path(tmp_path))
    with _client() as (client, session_factory):
        headers = _auth_headers(client)
        project_id = _project_id(client, headers)

        db = session_factory()
        try:
            project = db.get(Project, project_id)
            assert project is not None
            now = utc_now()
            historical = Scan(
                project=project,
                status="completed",
                scan_profile="safe",
                started_at=now - timedelta(hours=2),
                finished_at=now - timedelta(hours=1, minutes=55),
                assets_discovered=1,
                assets_scanned=1,
                findings_created=1,
                coverage_percent=100,
                risk_score=80,
            )
            latest = Scan(project=project, status="completed", scan_profile="safe", started_at=now - timedelta(minutes=10), finished_at=now)
            db.add_all([historical, latest])
            db.flush()
            canonical_asset = Asset(
                project=project,
                scan=latest,
                hostname="api.example.org",
                first_seen_at=now - timedelta(days=1),
                last_seen_at=now,
            )
            canonical_finding = Finding(
                project=project,
                scan=latest,
                asset=canonical_asset,
                title="Canonical row now belongs to latest scan",
                description="Current-state row.",
                severity="low",
                category="Current State",
                first_seen_at=now,
                last_seen_at=now,
            )
            manifest = ScanAssetResult(
                project_id=project.id,
                scan_id=historical.id,
                hostname="api.example.org",
                ip_addresses=["93.184.216.34"],
                source="certificate_transparency",
                discovery_status="active",
                scan_status="completed",
                risk_level="high",
                checks={"http": {"status": "completed"}},
                http_observations=[
                    {
                        "requested_url": "https://api.example.org",
                        "url": "https://api.example.org/login",
                        "status_code": 200,
                        "tls_verified": True,
                        "response_time_ms": 42,
                        "content_length": 512,
                        "redirect_chain": ["https://api.example.org/login"],
                        "error": None,
                    }
                ],
                port_observations=[
                    {
                        "port": 443,
                        "protocol": "tcp",
                        "status": "open",
                        "service_guess": "https",
                        "banner": None,
                        "error": None,
                    },
                    {
                        "port": 8080,
                        "protocol": "tcp",
                        "status": "closed",
                        "service_guess": "http-alt",
                        "banner": None,
                        "error": None,
                    },
                ],
                exposure_observations=[
                    {
                        "requested_url": "https://api.example.org/.env",
                        "url": "https://api.example.org/.env",
                        "path": "/.env",
                        "status_code": 404,
                        "content_type": "text/html",
                        "content_length": 128,
                        "tls_verified": True,
                        "response_time_ms": 35,
                        "redirect_chain": [],
                        "error": None,
                    }
                ],
                finding_observations=[
                    {
                        "title": "Historical critical exposure",
                        "description": "Preserved historical evidence.",
                        "severity": "critical",
                        "category": "Secret Exposure",
                        "status": "open",
                        "recommendation": "Rotate the exposed credential.",
                        "last_seen_at": historical.finished_at.isoformat(),
                    }
                ],
            )
            db.add_all([canonical_asset, canonical_finding, manifest])
            db.commit()
            historical_id = historical.id
        finally:
            db.close()

        response = client.post(
            f"/api/v1/projects/{project_id}/reports/excel",
            headers=headers,
            json={"scan_id": historical_id},
        )

        assert response.status_code == 201
        workbook = load_workbook(response.json()["file_path"])
        assert [row[0] for row in workbook["Assets"].iter_rows(min_row=2, values_only=True)] == ["api.example.org"]
        assert [row[2] for row in workbook["Findings"].iter_rows(min_row=2, values_only=True)] == [
            "Historical critical exposure"
        ]
        assert workbook["HTTP Endpoints"]["B2"].value == "https://api.example.org"
        assert [row[1] for row in workbook["Port Results"].iter_rows(min_row=2, values_only=True)] == [443, 8080]
        assert [row[1] for row in workbook["Open Ports"].iter_rows(min_row=2, values_only=True)] == [443]
        assert workbook["Exposure Checks"]["D2"].value == "/.env"

        pdf_response = client.post(
            f"/api/v1/projects/{project_id}/reports/pdf",
            headers=headers,
            json={"scan_id": historical_id},
        )
        assert pdf_response.status_code == 201
        pdf_path = Path(pdf_response.json()["file_path"])
        assert pdf_path.exists()
        assert pdf_path.stat().st_size > 0
